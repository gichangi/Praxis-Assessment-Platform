#!/usr/bin/env python3
"""
Convert the district assessment Excel template (e.g. 1308Assessing_template.xlsx)
into a form definition JSON that the backend, admin portal and mobile app all use.

Usage:
    python excel_to_template.py <input.xlsx> <output.json> [--version 1.0.0]

What it understands (the layout of the original sheet):
  * Rows 2-8, column A: the generic 0-6 maturity scale ("0 - Not yet started" ...).
  * Section 1: element rows between the "Element" header and the first
    "Total for next level dialogue" row. Element text in column E,
    guidance ("Conditions") in columns B and C, prompts in column N.
  * Section 2: after the first total. A row with text in F-L is a scale header;
    element rows below it use that scale until the next header.
  * Blank rows are skipped.

It prints a review report: every item, its group, scale, and any text it cleaned,
plus the two formula problems in the original workbook. A human should compare
that report with the sheet before the JSON is loaded into the database.
"""
import json
import re
import sys
from datetime import datetime, timezone

try:
    from openpyxl import load_workbook
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

SCORE_COLS = "FGHIJKL"  # values 0..6

# Typos in the source sheet, fixed on import and listed in the report.
TYPO_FIXES = {
    "settement": "settlement",
    "Distirct": "District",
    "distirct": "district",
    "avialable": "available",
    "adminisatration": "administration",
    "Interminant": "Intermittent",
    "Iniital": "Initial",
    "mechnaisms": "mechanisms",
    "donr": "done",
    "compre-henisvely": "comprehensively",
}

# Words the sheet hyphenated only to fit narrow columns.
DEHYPHENATE = {
    "Humani-tarian": "Humanitarian",
    "Compre-hensive": "Comprehensive",
    "engage-ment": "engagement",
    "consol-idation": "consolidation",
    "Acknow-ledge": "Acknowledge",
    "inter-ventions": "interventions",
    "imple-mentation": "implementation",
}

# Short labels for the generic scale (the full sheet text becomes the description).
MATURITY_LABELS = {
    0: "Not yet started",
    1: "Awareness & dialogue",
    2: "Emerging practice",
    3: "Practicing",
    4: "Working",
    5: "Established",
    6: "Fully functioning",
}

# Section 1 groups inferred from the sheet's structure (confirm with the client).
SECTION1_GROUPS = [
    ("council_stability", "Council formation & stability", 10, 17),
    ("participation_resources", "Participation & resources", 18, 21),
    ("technical_services", "Technical & service capacity", 22, 33),
    ("effectiveness_systems", "Government effectiveness & systems", 34, 49),
]


def clean(text, report, where):
    if text is None:
        return None
    original = str(text)
    t = re.sub(r"\s+", " ", original).strip()
    for bad, good in {**TYPO_FIXES, **DEHYPHENATE}.items():
        t = t.replace(bad, good)
    t = t.replace("’", "'")
    if t != re.sub(r"\s+", " ", original).strip():
        report.append(f"  cleaned {where}: '{original.strip()[:60]}' -> '{t[:60]}'")
    return t or None


def slug(text, used):
    words = re.sub(r"[^a-z0-9 ]", " ", text.lower()).split()
    stop = {"the", "of", "and", "a", "an", "to", "in", "for", "as", "is", "by", "with", "on", "what", "how", "are", "or", "there", "be", "who", "which"}
    words = [w for w in words if w not in stop][:4] or ["item"]
    base = "_".join(words)
    key, n = base, 2
    while key in used:
        key, n = f"{base}_{n}", n + 1
    used.add(key)
    return key


def i18n(text):
    return {"en": text} if text else None


def group_for(row):
    for key, _title, start, end in SECTION1_GROUPS:
        if start <= row <= end:
            return key
    return SECTION1_GROUPS[-1][0]


def convert(path, version):
    ws = load_workbook(path, data_only=False).active
    report, used_keys = [], set()
    cell = lambda r, c: ws[f"{c}{r}"].value

    # 1. Generic maturity scale
    levels = []
    for r in range(2, 9):
        raw = clean(cell(r, "A"), report, f"A{r}")
        m = re.match(r"(\d)\s*-\s*(.*)", raw or "")
        if m:
            v = int(m.group(1))
            levels.append({"value": v, "label": i18n(MATURITY_LABELS.get(v, m.group(2))),
                           "description": i18n(m.group(2))})
    scales = [{"key": "maturity_0_6", "title": i18n("District development maturity"), "levels": levels}]

    # 2. Find section boundaries
    totals = [r for r in range(1, ws.max_row + 1)
              if isinstance(cell(r, "E"), str) and "Total for next level" in cell(r, "E")]
    header = next(r for r in range(1, ws.max_row + 1) if cell(r, "E") == "Element")
    s1_end, s2_end = totals[0], totals[1] if len(totals) > 1 else ws.max_row + 1

    # 3. Section 1 items
    s1_items = []
    for r in range(header + 2, s1_end):
        title = clean(cell(r, "E"), report, f"E{r}")
        if not title:
            continue
        guidance = " ".join(filter(None, [clean(cell(r, "B"), report, f"B{r}"), clean(cell(r, "C"), report, f"C{r}")]))
        prompt = clean(cell(r, "N"), report, f"N{r}")
        has_formula = isinstance(cell(r, "M"), str) and cell(r, "M").startswith("=")
        if not has_formula:
            report.append(f"  NOTE row {r}: no total formula in the Excel, so it was never counted. Included here.")
        s1_items.append({
            "key": slug(title, used_keys), "type": "scale", "scale": "maturity_0_6",
            "group": group_for(r), "title": i18n(title), "guidance": i18n(guidance or None),
            "prompt": i18n(prompt), "required": False, "weight": 1, "include_in_total": True,
            "source_row": r,
        })

    # 4. Section 2 items with per-criterion scales
    s2_items, current_scale, scale_n = [], None, 0
    section2_title, section2_guidance = None, []
    for r in range(s1_end + 1, s2_end):
        labels = [clean(cell(r, c), report, f"{c}{r}") for c in SCORE_COLS]
        e_text = clean(cell(r, "E"), report, f"E{r}")
        b_text = clean(cell(r, "B"), report, f"B{r}")
        if b_text:
            section2_guidance.append(b_text)
        if all(labels):
            scale_n += 1
            current_scale = f"logistics_scale_{scale_n}"
            scales.append({"key": current_scale, "title": i18n(f"Logistics scale {scale_n}"),
                           "levels": [{"value": i, "label": i18n(l), "description": None} for i, l in enumerate(labels)]})
            if e_text and not section2_title:
                section2_title = e_text
            continue
        if e_text and current_scale:
            s2_items.append({
                "key": slug(e_text, used_keys), "type": "scale", "scale": current_scale,
                "group": "logistics", "title": i18n(e_text), "guidance": None,
                "prompt": i18n(clean(cell(r, "N"), report, f"N{r}")), "required": False,
                "weight": 1, "include_in_total": True, "source_row": r,
            })

    formula = cell(s2_end, "M") if s2_end <= ws.max_row else None
    if isinstance(formula, str) and "54" in formula:
        report.append(f"  NOTE Section 2 total in the Excel is {formula}: it skips the Accessibility row. "
                      "The app computes totals from the definition, so every item counts.")

    title = clean(cell(1, "E"), report, "E1")
    definition = {
        "schema_version": 1,
        "form_key": "district_governance_assessment",
        "version": version,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "default_language": "en",
        "languages": ["en", "so", "ar"],
        "title": i18n(title),
        "scoring": {"method": "sum", "exclude_not_assessed": True, "show_percentage": True},
        "scales": scales,
        "sections": [
            {"key": "governance", "title": i18n("Governance elements"),
             "description": i18n("Has the District developed these elements?"),
             "groups": [{"key": k, "title": i18n(t)} for k, t, _s, _e in SECTION1_GROUPS],
             "items": s1_items},
            {"key": "logistics", "title": i18n(section2_title or "Initial logistical engagement considerations"),
             "description": i18n(" ".join(section2_guidance) or None),
             "groups": [{"key": "logistics", "title": i18n("Logistics & engagement")}],
             "items": s2_items},
        ],
    }
    return definition, report


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    version = "1.0.0"
    if "--version" in sys.argv:
        version = sys.argv[sys.argv.index("--version") + 1]
    definition, report = convert(sys.argv[1], version)
    with open(sys.argv[2], "w", encoding="utf-8") as f:
        json.dump(definition, f, ensure_ascii=False, indent=2)

    print(f"Wrote {sys.argv[2]}")
    for s in definition["sections"]:
        print(f"\nSection '{s['key']}': {len(s['items'])} items")
        for it in s["items"]:
            print(f"  row {it['source_row']:>3}  [{it['group']}] {it['key']:<40} scale={it['scale']}")
    print(f"\nScales: {', '.join(sc['key'] for sc in definition['scales'])}")
    print("\nReview notes:")
    print("\n".join(report) if report else "  none")


if __name__ == "__main__":
    main()
